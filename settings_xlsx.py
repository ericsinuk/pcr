#!/usr/bin/env python3
"""Generate the master settings workbook from the app source (app.js + airfields.js).

The workbook is the editing interface for all app data: fleet, chart values,
MTW/MLW reference weights, and the airfield database. Edit it, hand it back,
and the app data files get regenerated from it.

Usage: python3 settings_xlsx.py [output.xlsx]
"""
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

app_src = open("app.js").read()
af_src = open("airfields.js").read()

ver = re.search(r'APP_VERSION = "([^"]+)"', app_src).group(1)
ver_date = re.search(r'APP_VERSION_DATE = "([^"]+)"', app_src).group(1)

fleet = re.findall(r'reg:\s*"([\w\-]+)",\s*type:\s*"([\w\-]+)",\s*minWt:\s*([\d.]+)', app_src)
fleet = [(r, t, float(w)) for r, t, w in fleet]

types_block = app_src[app_src.index("const AIRCRAFT_TYPES"):app_src.index("const TYPE_BADGE")]
chart = {}
for tm in re.finditer(r'"([\w\-]+)":\s*\{\s*maxTaxi:\s*(\d+),(.*?)\n  \},', types_block, re.S):
    typ, maxtaxi, body = tm.group(1), int(tm.group(2)), tm.group(3)
    pcr_part, pcn_part = body.split("PCN:")
    parse = lambda seg: {m.group(1): (int(m.group(2)), int(m.group(3)))
                         for m in re.finditer(r'(\w\w):\{max:(\d+),min:(\d+)\}', seg)}
    chart[typ] = {"maxTaxi": maxtaxi, "PCR": parse(pcr_part), "PCN": parse(pcn_part)}

refs_block = app_src[app_src.index("const REF_WEIGHTS"):app_src.index("const CHK_TYPES")]
refs = {}
for m in re.finditer(r'(MTW|MLW):\s*\{([^}]*)\}', refs_block):
    refs[m.group(1)] = dict(re.findall(r'"([\w\-]+)":\s*(\d+)', m.group(2)))

af_json = af_src[af_src.index("const AIRFIELDS = ") + len("const AIRFIELDS = "):]
airfields = json.loads(af_json[:af_json.rindex("]") + 1])

# ---- workbook ----
F = lambda **kw: Font(name="Arial", size=10, **kw)
HDR = PatternFill("solid", start_color="FFCC00")
YEL = PatternFill("solid", start_color="FFFF00")
thin = Side(style="thin", color="C0C0C0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def head(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = F(bold=True); cell.fill = HDR; cell.border = BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

# README
ws = wb.active; ws.title = "README"
lines = [
 ("These are the numbers used in the PCN/PCR calculator.", False),
 ("if you have any questions, please email efb.helpdesk@dhl.com", False),
]
for i, (txt, bold) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=txt).font = F(bold=bold)
ws.column_dimensions["A"].width = 110

# Fleet
ws = wb.create_sheet("Fleet")
head(ws, 1, ["Registration", "Type", "Min Weight (kg)"], (14, 12, 16))
dv = DataValidation(type="list", formula1='"757-200,767-300ER,777F"', allow_blank=True)
ws.add_data_validation(dv); dv.add("B2:B100")
for r, (reg, typ, mw) in enumerate(fleet, 2):
    ws.cell(row=r, column=1, value=reg).border = BOX
    ws.cell(row=r, column=2, value=typ).border = BOX
    c = ws.cell(row=r, column=3, value=mw); c.border = BOX; c.fill = YEL; c.number_format = "#,##0.00"
    for col in range(1, 4): ws.cell(row=r, column=col).font = F()

# Aircraft Data
ws = wb.create_sheet("Aircraft Data")
head(ws, 1, ["Type", "Max Taxi (kg)", "Pavement", "Subgrade", "Code",
             "PCR @MaxTaxi", "PCR @MinWt", "PCN @MaxTaxi", "PCN @MinWt", "Key"],
     (12, 13, 10, 9, 7, 12, 12, 12, 12, 14))
r = 2
for typ in chart:
    for code in ("RA", "RB", "RC", "RD", "FA", "FB", "FC", "FD"):
        pcr, pcn = chart[typ]["PCR"][code], chart[typ]["PCN"][code]
        vals = [typ, chart[typ]["maxTaxi"], "Rigid" if code[0] == "R" else "Flexible",
                code[1], code, pcr[0], pcr[1], pcn[0], pcn[1]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = F(); cell.border = BOX
            if c in (6, 7, 8, 9): cell.fill = YEL
        ws.cell(row=r, column=2).number_format = "#,##0"
        k = ws.cell(row=r, column=10, value=f"=A{r}&E{r}"); k.font = F(color="808080"); k.border = BOX
        r += 1

# Ref Weights
ws = wb.create_sheet("Ref Weights")
head(ws, 1, ["Type", "MTW (kg)", "MLW (kg)"], (12, 14, 14))
for r, typ in enumerate(chart, 2):
    ws.cell(row=r, column=1, value=typ).font = F(); ws.cell(row=r, column=1).border = BOX
    for c, key in ((2, "MTW"), (3, "MLW")):
        cell = ws.cell(row=r, column=c, value=int(refs[key][typ]))
        cell.font = F(); cell.border = BOX; cell.fill = YEL; cell.number_format = "#,##0"

# Airfields — one row per runway
ws = wb.create_sheet("Airfields")
head(ws, 1, ["ICAO", "IATA", "Airfield", "Ops Type", "Status", "System",
             "RWY", "Length (ft)", "Width (ft)", "PCN/PCR"],
     (8, 7, 26, 24, 24, 8, 6, 11, 10, 24))
dv2 = DataValidation(type="list", formula1='"PCR,PCN"', allow_blank=True)
ws.add_data_validation(dv2); dv2.add("F2:F3000")
r = 2
for icao, iata, name, ops, status, system, rwys in airfields:
    for rwy, ln, w, pcn in rwys:
        for c, v in enumerate([icao, iata, name, ops, status, system, rwy, ln or "", w or "", pcn], 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = F(); cell.border = BOX
        ws.cell(row=r, column=10).fill = YEL
        r += 1
ws.freeze_panes = "A2"

out = sys.argv[1] if len(sys.argv) > 1 else "PCN_PCR_Settings.xlsx"
wb.save(out)
print(f"saved {out} — Ver {ver} ({ver_date}): {len(fleet)} aircraft, "
      f"{sum(1 for _ in chart)} types, {len(airfields)} airfields")
